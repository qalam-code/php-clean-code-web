# -*- coding: utf-8 -*-
"""
Genere Documentation_phpCleanCode.docx a partir des fichiers Markdown de docs/.

    python generer_doc.py [dossier-docs] [fichier-sortie]

Les fichiers Markdown sont la reference : ils vivent a cote du code et se
mettent a jour avec lui. Ce script n'est qu'un convertisseur, pour produire une
version imprimable ou transmissible. NE CORRIGEZ JAMAIS LE .docx A LA MAIN :
corrigez le .md et relancez ce script, sinon les deux divergent en une semaine.

Dependance : python-docx  (pip install python-docx)
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

VERSION = "1.0.0"
TITRE = "phpCleanCode"
SOUS_TITRE = "Socle Clean Architecture pour API PHP"
ORGANISATION = "Flexeau Afrique"

# Ordre de composition. Un fichier absent est signale, pas ignore en silence.
ORDRE = [
    "README.md",
    "01-principes.md",
    "02-domaine.md",
    "03-infrastructure.md",
    "04-http.md",
    "05-presentation.md",
    "06-composition.md",
    "07-tests.md",
    "08-outils.md",
    "09-squelette.md",
    "10-recettes.md",
    "11-pieges.md",
]

BLEU = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = RGBColor(0x55, 0x55, 0x55)
ROUGE = RGBColor(0x9B, 0x2C, 0x2C)
FOND_CODE = "F4F4F4"
FOND_ENTETE = "1F3A5F"
FOND_ALERTE = "FFF4E5"
FOND_LIGNE = "FAFAFA"


# ---------------------------------------------------------------------------
# Helpers XML bas niveau
# ---------------------------------------------------------------------------

def fond_cellule(cellule, couleur_hex):
    tcPr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur_hex)
    tcPr.append(shd)


def marges_cellule(cellule, haut=60, bas=60, gauche=110, droite=110):
    tcPr = cellule._tc.get_or_add_tcPr()
    marges = OxmlElement("w:tcMar")
    for nom, valeur in (("top", haut), ("bottom", bas), ("start", gauche), ("end", droite)):
        n = OxmlElement("w:" + nom)
        n.set(qn("w:w"), str(valeur))
        n.set(qn("w:type"), "dxa")
        marges.append(n)
    tcPr.append(marges)


def bordure_basse(paragraphe, couleur="CCCCCC", taille=6):
    pPr = paragraphe._p.get_or_add_pPr()
    bordures = OxmlElement("w:pBdr")
    bas = OxmlElement("w:bottom")
    bas.set(qn("w:val"), "single")
    bas.set(qn("w:sz"), str(taille))
    bas.set(qn("w:space"), "4")
    bas.set(qn("w:color"), couleur)
    bordures.append(bas)
    pPr.append(bordures)


def champ(paragraphe, instruction):
    """Insere un champ Word (TOC, PAGE...)."""
    run = paragraphe.add_run()
    debut = OxmlElement("w:fldChar")
    debut.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separe = OxmlElement("w:fldChar")
    separe.set(qn("w:fldCharType"), "separate")
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for n in (debut, instr, separe, fin):
        run._r.append(n)


# ---------------------------------------------------------------------------
# Rendu du texte en ligne : **gras**, `code`, *italique*
# ---------------------------------------------------------------------------

JETON = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def texte_enrichi(paragraphe, texte, base_gras=False):
    for morceau in JETON.split(texte):
        if not morceau:
            continue
        if morceau.startswith("**") and morceau.endswith("**"):
            # Un code en gras reste du gras : on retire les accents graves
            # plutot que de les laisser apparaitre dans le texte.
            run = paragraphe.add_run(morceau[2:-2].replace("`", ""))
            run.bold = True
        elif morceau.startswith("`") and morceau.endswith("`"):
            run = paragraphe.add_run(morceau[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = ROUGE
        elif morceau.startswith("*") and morceau.endswith("*") and len(morceau) > 2:
            run = paragraphe.add_run(morceau[1:-1])
            run.italic = True
        elif morceau.startswith("["):
            # [libelle](cible) : on ne garde que le libelle, les liens entre
            # fichiers .md n'ont pas de sens dans un document unique.
            run = paragraphe.add_run(re.sub(r"^\[([^\]]+)\].*$", r"\1", morceau))
            run.italic = True
            run.font.color.rgb = GRIS
        else:
            run = paragraphe.add_run(morceau)
        if base_gras:
            run.bold = True


# ---------------------------------------------------------------------------
# Blocs
# ---------------------------------------------------------------------------

def bloc_code(doc, lignes):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cellule = table.rows[0].cells[0]
    fond_cellule(cellule, FOND_CODE)
    marges_cellule(cellule, 120, 120, 140, 140)
    cellule.text = ""
    premier = True
    for ligne in lignes:
        p = cellule.paragraphs[0] if premier else cellule.add_paragraph()
        premier = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(ligne if ligne else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def bloc_alerte(doc, lignes):
    # Les lignes vides de tete et de queue ne servent a rien dans un encadre.
    while lignes and not lignes[0].strip():
        lignes.pop(0)
    while lignes and not lignes[-1].strip():
        lignes.pop()
    if not lignes:
        return

    table = doc.add_table(rows=1, cols=1)
    cellule = table.rows[0].cells[0]
    fond_cellule(cellule, FOND_ALERTE)
    marges_cellule(cellule, 120, 120, 160, 160)
    cellule.text = ""

    premier = [True]

    def paragraphe():
        if premier[0]:
            premier[0] = False
            return cellule.paragraphs[0]
        return cellule.add_paragraph()

    i = 0
    tampon = []

    def vider():
        if not tampon:
            return
        p = paragraphe()
        p.paragraph_format.space_after = Pt(4)
        texte_enrichi(p, " ".join(tampon))
        del tampon[:]

    while i < len(lignes):
        ligne = lignes[i]
        # Un bloc de code peut vivre dans une citation : on le rend en
        # monospace, sans imbriquer un tableau dans un tableau.
        if ligne.strip().startswith("```"):
            vider()
            i += 1
            while i < len(lignes) and not lignes[i].strip().startswith("```"):
                p = paragraphe()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(lignes[i] if lignes[i] else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                i += 1
            i += 1
            continue
        if not ligne.strip():
            vider()
            i += 1
            continue
        tampon.append(ligne.strip())
        i += 1
    vider()

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def bloc_table(doc, lignes):
    """lignes : liste de lignes Markdown '| a | b |', separateur inclus."""
    donnees = []
    for ligne in lignes:
        if re.match(r"^\|[\s:\-|]+\|$", ligne.strip()):
            continue
        cellules = [c.strip() for c in ligne.strip().strip("|").split("|")]
        donnees.append(cellules)
    if not donnees:
        return

    colonnes = max(len(l) for l in donnees)
    table = doc.add_table(rows=0, cols=colonnes)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Largeurs proportionnelles au contenu, bornees pour qu'aucune colonne
    # ne devienne illisible. Sans cela, une colonne de numeros occupe le
    # tiers de la page.
    largeurs = []
    for j in range(colonnes):
        longueurs = [len(l[j]) for l in donnees if j < len(l)]
        largeurs.append(max(6, min(60, max(longueurs) if longueurs else 6)))
    total = float(sum(largeurs))
    utile = Cm(16.2)
    table.autofit = False
    for j, part in enumerate(largeurs):
        largeurs[j] = int(utile * (part / total))

    # Word et LibreOffice n'obeissent aux largeurs que si elles sont posees
    # A LA FOIS sur la grille du tableau et sur chaque cellule, en disposition
    # fixe. Une seule des trois, et la colonne revient a sa largeur calculee.
    tblPr = table._tbl.tblPr
    disposition = OxmlElement("w:tblLayout")
    disposition.set(qn("w:type"), "fixed")
    tblPr.append(disposition)
    for j, largeur in enumerate(largeurs):
        table.columns[j].width = largeur

    for index, ligne in enumerate(donnees):
        cellules = table.add_row().cells
        for j in range(colonnes):
            contenu = ligne[j] if j < len(ligne) else ""
            cellule = cellules[j]
            cellule.width = largeurs[j]
            cellule.text = ""
            marges_cellule(cellule)
            p = cellule.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index == 0:
                fond_cellule(cellule, FOND_ENTETE)
                run = p.add_run(re.sub(r"[*`]", "", contenu))
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                if index % 2 == 0:
                    fond_cellule(cellule, FOND_LIGNE)
                texte_enrichi(p, contenu)
                for run in p.runs:
                    if run.font.size is None:
                        run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ---------------------------------------------------------------------------
# Analyse d'un fichier Markdown
# ---------------------------------------------------------------------------

def rendre_markdown(doc, texte):
    lignes = texte.split("\n")
    i = 0
    n = len(lignes)

    while i < n:
        ligne = lignes[i]
        nu = ligne.strip()

        # --- bloc de code ---
        if nu.startswith("```"):
            i += 1
            code = []
            while i < n and not lignes[i].strip().startswith("```"):
                code.append(lignes[i])
                i += 1
            i += 1
            bloc_code(doc, code)
            continue

        # --- citation / alerte ---
        if nu.startswith(">"):
            contenu = []
            while i < n and lignes[i].strip().startswith(">"):
                contenu.append(lignes[i].strip().lstrip(">").strip())
                i += 1
            bloc_alerte(doc, [c for c in contenu if c or contenu.index(c) not in (0, len(contenu) - 1)])
            continue

        # --- tableau ---
        if nu.startswith("|"):
            bloc = []
            while i < n and lignes[i].strip().startswith("|"):
                bloc.append(lignes[i])
                i += 1
            bloc_table(doc, bloc)
            continue

        # --- separateur horizontal ---
        if nu in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            bordure_basse(p)
            i += 1
            continue

        # --- titres ---
        if nu.startswith("#"):
            niveau = len(nu) - len(nu.lstrip("#"))
            # Un titre ne porte pas de mise en forme : on retire les marqueurs
            # Markdown plutot que de les afficher tels quels.
            titre = re.sub(r"[`*]", "", nu.lstrip("#").strip())
            if niveau == 1:
                p = doc.add_heading(titre, level=1)
            elif niveau == 2:
                p = doc.add_heading(titre, level=2)
            else:
                p = doc.add_heading(titre, level=3)
            for run in p.runs:
                run.font.color.rgb = BLEU
            i += 1
            continue

        # --- listes ---
        puce = re.match(r"^[-*] ", nu)
        numero = re.match(r"^\d+\. ", nu)
        if puce or numero:
            premiere = nu[2:] if puce else re.sub(r"^\d+\.\s*", "", nu)
            i += 1
            # Un element de liste qui deborde sur la ligne suivante reste le
            # meme element : sans cela, la suite devient un paragraphe orphelin.
            suite = [premiere]
            while i < n and lignes[i].strip() and not re.match(
                r"^(#|```|\||>|[-*] |\d+\. |---$|\*\*\*$|___$)", lignes[i].strip()
            ):
                suite.append(lignes[i].strip())
                i += 1
            p = doc.add_paragraph(style="List Bullet" if puce else "List Number")
            p.paragraph_format.space_after = Pt(3)
            texte_enrichi(p, " ".join(suite))
            continue

        # --- ligne vide ---
        if not nu:
            i += 1
            continue

        # --- paragraphe : on recolle les lignes jusqu'au prochain blanc ---
        bloc = []
        while i < n and lignes[i].strip() and not re.match(
            r"^(#|```|\||>|[-*] |\d+\. |---$|\*\*\*$|___$)", lignes[i].strip()
        ):
            bloc.append(lignes[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        texte_enrichi(p, " ".join(bloc))


# ---------------------------------------------------------------------------
# Page de garde et sommaire
# ---------------------------------------------------------------------------

def page_de_garde(doc):
    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITRE)
    run.bold = True
    run.font.size = Pt(40)
    run.font.color.rgb = BLEU

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(SOUS_TITRE)
    run.font.size = Pt(15)
    run.font.color.rgb = GRIS

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Documentation technique d\u00e9taill\u00e9e")
    run.font.size = Pt(12)
    run.italic = True

    for _ in range(8):
        doc.add_paragraph()

    for ligne, taille in (
        (ORGANISATION, 12),
        ("PHP 7.0 minimum   |   aucune d\u00e9pendance", 10),
        ("Version du document " + VERSION + "   |   24 septembre 2026", 10),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(ligne)
        run.font.size = Pt(taille)
        run.font.color.rgb = GRIS

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def sommaire(doc):
    p = doc.add_heading("Sommaire", level=1)
    for run in p.runs:
        run.font.color.rgb = BLEU

    p = doc.add_paragraph()
    run = p.add_run(
        "Ce sommaire est un champ Word. Pour le remplir : ouvrir le document, "
        "clic droit dessus, Â« Mettre Ã  jour les champs Â»."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GRIS

    champ(doc.add_paragraph(), r'TOC \o "1-2" \h \z \u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def pied_de_page(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(TITRE + "  â€”  ")
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS
        champ(p, "PAGE")
        for run in p.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = GRIS


# ---------------------------------------------------------------------------

def main():
    dossier = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
    sortie = sys.argv[2] if len(sys.argv) > 2 else "Documentation_phpCleanCode.docx"

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    page_de_garde(doc)
    sommaire(doc)

    manquants = []
    for index, nom in enumerate(ORDRE):
        chemin = os.path.join(dossier, nom)
        if not os.path.isfile(chemin):
            manquants.append(nom)
            continue
        with open(chemin, "r", encoding="utf-8") as f:
            contenu = f.read()
        if index > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        rendre_markdown(doc, contenu)
        print("  + " + nom)

    pied_de_page(doc)
    doc.save(sortie)

    print("")
    print(sortie + " genere.")
    if manquants:
        print("ATTENTION, fichiers absents : " + ", ".join(manquants))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

